#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按火山引擎内部报价模板生成报价单 xlsx。

用法:
    python3 build_quote.py spec.json

spec.json 格式:
{
  "base": "/workspace/.uploads/xxx-报价参考.xlsx",   // 内部报价模板（样式来源）
  "out":  "/workspace/火山引擎报价单-北京-xxx.xlsx",  // 输出文件
  "report_date": "2026-10-16",                       // 可选，默认今天
  "items": [
    {
      "product": "云服务器",                          // 产品名
      "region":  "北京",                              // 地域
      "config":  "实例规格：ecs.c4i.xlarge（4vCPU/8GiB）\n系统盘：...",  // 配置列，\n 换行
      "config_fee": 3328.08,                          // 采购界面「配置费用」
      "saved":      1687.92                           // 采购界面「已省费用」
    }
  ]
}

口径: 官网刊例价 = 配置费用; 官网目录价 = 配置费用 + 已省费用; 总计 = SUM。
"""
import json
import os
import re
import shutil
import sys
import zipfile
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.styles import Alignment


def _find_row(ws, text, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is not None and text in str(v):
            return r
    return None


def build(spec):
    base, out = spec["base"], spec["out"]
    items = spec["items"]
    if not items:
        raise SystemExit("items 为空")

    shutil.copy(base, out)
    wb = load_workbook(out)
    ws = wb.active

    header_row = _find_row(ws, "序号") or 3
    total_row = _find_row(ws, "总计", header_row + 1) or ws.max_row
    first_prod = header_row + 1

    title_merges = [str(r) for r in ws.merged_cells.ranges if r.min_row == 1]

    # 解除全部合并，使所有单元格可写
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))

    def snap(row):
        return {c: copy(ws.cell(row, i)._style) for i, c in enumerate("ABCDEFG", 1)}

    prod_style = snap(first_prod)
    total_style = snap(total_row)
    total_note = ws.cell(total_row, 7).value
    total_height = ws.row_dimensions[total_row].height or 35
    numfmt_e = ws.cell(first_prod, 5).number_format
    numfmt_f = ws.cell(first_prod, 6).number_format
    font_size = ws.cell(first_prod, 4).font.size or 12

    # 清空第 first_prod 行及以后
    for r in range(first_prod, ws.max_row + 1):
        for i in range(1, 8):
            ws.cell(r, i).value = None

    n = len(items)
    new_total = first_prod + n

    # 产品行
    for idx, it in enumerate(items):
        r = first_prod + idx
        for c in "ABCDEFG":
            ws[f"{c}{r}"]._style = copy(prod_style[c])
        ws.cell(r, 1).value = idx + 1
        ws.cell(r, 2).value = it.get("product", "")
        ws.cell(r, 3).value = it.get("region", "")
        ws.cell(r, 4).value = it.get("config", "")
        ws.cell(r, 5).value = "=%s+%s" % (it["config_fee"], it["saved"])
        ws.cell(r, 6).value = it["config_fee"]
        ws.cell(r, 7).value = None
        al = copy(ws.cell(r, 4).alignment)
        al.wrapText = True
        ws.cell(r, 4).alignment = al
        ws.cell(r, 5).number_format = numfmt_e
        ws.cell(r, 6).number_format = numfmt_f
        lines = str(it.get("config", "")).count("\n") + 1
        ws.row_dimensions[r].height = max(45.0, round(lines * font_size * 1.45, 1))

    # 总计行
    for c in "ABCDEFG":
        ws[f"{c}{new_total}"]._style = copy(total_style[c])
    ws.cell(new_total, 1).value = "总计"
    ws.cell(new_total, 5).value = "=SUM(E%d:E%d)" % (first_prod, first_prod + n - 1)
    ws.cell(new_total, 6).value = "=SUM(F%d:F%d)" % (first_prod, first_prod + n - 1)
    ws.cell(new_total, 7).value = total_note
    ws.cell(new_total, 5).number_format = numfmt_e
    ws.cell(new_total, 6).number_format = numfmt_f
    ws.cell(new_total, 7).alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    ws.row_dimensions[new_total].height = total_height

    # 删除多余行
    if ws.max_row > new_total:
        ws.delete_rows(new_total + 1, ws.max_row - new_total)
    for r in list(ws.row_dimensions.keys()):
        if r > new_total:
            del ws.row_dimensions[r]

    # 重设合并
    for rng in title_merges:
        try:
            ws.merge_cells(rng)
        except Exception:
            pass
    ws.merge_cells(f"A{new_total}:D{new_total}")

    wb.save(out)

    # 注入公式缓存值（未重算环境也能正确显示）
    rd = spec.get("report_date")
    d = date.fromisoformat(rd) if rd else date.today()
    serial = (d - date(1899, 12, 30)).days
    cached = {"G2": str(serial)}
    fee_sum = saved_sum = 0.0
    for idx, it in enumerate(items):
        r = first_prod + idx
        cached[f"E{r}"] = _num(it["config_fee"] + it["saved"])
        fee_sum += it["config_fee"]
        saved_sum += it["saved"]
    cached[f"E{new_total}"] = _num(fee_sum + saved_sum)
    cached[f"F{new_total}"] = _num(fee_sum)

    zin = zipfile.ZipFile(out)
    parts = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    sheet = [n for n in parts if re.match(r"xl/worksheets/sheet\d+\.xml$", n)][0]
    xml = parts[sheet].decode("utf-8")
    for ref, val in cached.items():
        xml, k = re.subn(
            r'(<c r="%s"[^>]*>(?:<f>[^<]*</f>))' % ref,
            lambda m: m.group(1) + "<v>%s</v>" % val,
            xml, count=1,
        )
        if not k:
            print("WARN: 未找到公式单元格 %s" % ref)
    parts[sheet] = xml.encode("utf-8")
    z = zipfile.ZipFile(out + ".tmp", "w", zipfile.ZIP_DEFLATED)
    for name in parts:
        z.writestr(name, parts[name])
    z.close()
    os.replace(out + ".tmp", out)

    print("OK ->", out)
    print("产品行 %d 行，总计行第 %d 行" % (n, new_total))
    for ref in sorted(cached, key=lambda x: (len(x), x)):
        print("  %s = %s" % (ref, cached[ref]))


def _num(x):
    x = round(float(x), 2)
    return str(int(x)) if abs(x - round(x)) < 1e-9 else str(x)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        raise SystemExit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f))
